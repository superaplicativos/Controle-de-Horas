import sys, os
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Fechamento Setembro 2026"

# ===== CORES =====
VERDE_ESCURO = "0A1410"
VERDE = "059669"
DOURADO = "FBBF24"
CINZA_CLARO = "F3F4F6"
CINZA = "6B7280"
VERMELHO = "EF4444"
BRANCO = "FFFFFF"

# ===== BORDAS =====
borda_fina = Side(style='thin', color='D1D5DB')
borda_titulo = Side(style='medium', color=VERDE_ESCURO)
borda = Border(left=borda_fina, right=borda_fina, top=borda_fina, bottom=borda_fina)

# ===== FONTES =====
font_titulo = Font(name='Calibri', size=18, bold=True, color=BRANCO)
font_subtitulo = Font(name='Calibri', size=12, color=BRANCO)
font_professor = Font(name='Calibri', size=14, bold=True, color=VERDE_ESCURO)
font_header = Font(name='Calibri', size=11, bold=True, color=BRANCO)
font_normal = Font(name='Calibri', size=11, color='1F2937')
font_subtotal = Font(name='Calibri', size=11, bold=True, color=VERDE_ESCURO)
font_total = Font(name='Calibri', size=14, bold=True, color=DOURADO)
font_rodape = Font(name='Calibri', size=9, color=CINZA, italic=True)

# ===== FILLS =====
fill_titulo = PatternFill(start_color=VERDE_ESCURO, end_color=VERDE_ESCURO, fill_type='solid')
fill_header = PatternFill(start_color=VERDE, end_color=VERDE, fill_type='solid')
fill_subtotal = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')
fill_total = PatternFill(start_color=VERDE_ESCURO, end_color=VERDE_ESCURO, fill_type='solid')
fill_zebra = PatternFill(start_color=CINZA_CLARO, end_color=CINZA_CLARO, fill_type='solid')
fill_falta = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

# ===== ALINHAMENTO =====
align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

# ===== LARGURA COLUNAS =====
ws.column_dimensions['A'].width = 5
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 10
ws.column_dimensions['E'].width = 10
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 35
ws.column_dimensions['H'].width = 14

linha = 1

# ===== CABEÇALHO =====
ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = "CONTROLE DE AULAS"
ws[f'A{linha}'].font = font_titulo
ws[f'A{linha}'].fill = fill_titulo
ws[f'A{linha}'].alignment = align_center
ws.row_dimensions[linha].height = 35
linha += 1

ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = "Fechamento Mensal — Setembro 2026"
ws[f'A{linha}'].font = font_subtitulo
ws[f'A{linha}'].fill = fill_titulo
ws[f'A{linha}'].alignment = align_center
ws.row_dimensions[linha].height = 25
linha += 1

ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = "Professor: Guilherme Miranda  |  Valor/hora: R$ 35,00"
ws[f'A{linha}'].font = font_professor
ws[f'A{linha}'].alignment = align_center
ws.row_dimensions[linha].height = 22
linha += 2

# ===== DADOS =====
dados = [
    # (aluno, dia, dia_semana, horario, duracao, status, conteudo, valor)
    # Thiago Patrick Bellini
    ("Thiago Patrick Bellini", "08/09", "Terça", "", 1, "Falta", "Excel / Mkt Digital", 35.00),
    ("Thiago Patrick Bellini", "15/09", "Terça", "", 2, "Presença", "Excel / Mkt Digital", 70.00),
    ("Thiago Patrick Bellini", "22/09", "Terça", "", 1, "Falta", "Excel / Mkt Digital", 35.00),
    ("Thiago Patrick Bellini", "24/09", "Quinta", "", 2, "Presença", "Excel / Mkt Digital", 70.00),
    ("Thiago Patrick Bellini", "29/09", "Terça", "", 1, "Falta", "Excel / Mkt Digital", 35.00),
    # Marcelo
    ("Marcelo Dercatshoff Pereira", "02/09", "Quarta", "", 2, "Presença", "Marketing Digital", 70.00),
    ("Marcelo Dercatshoff Pereira", "09/09", "Quarta", "", 2, "Presença", "Marketing Digital", 70.00),
    ("Marcelo Dercatshoff Pereira", "16/09", "Quarta", "", 2, "Presença", "Marketing Digital", 70.00),
    ("Marcelo Dercatshoff Pereira", "30/09", "Quarta", "", 2, "Presença", "Marketing Digital", 70.00),
    # Pedro Henrique
    ("Pedro Henrique Russo Nardi", "11/09", "Sexta", "", 2, "Presença", "IA / Lógica de Programação", 70.00),
    ("Pedro Henrique Russo Nardi", "18/09", "Sexta", "", 2, "Presença", "IA / Lógica de Programação", 70.00),
    ("Pedro Henrique Russo Nardi", "25/09", "Sexta", "", 2, "Presença", "IA / Lógica de Programação", 70.00),
    # Maria José
    ("Maria José Junqueira da Silva", "05/09", "Sábado", "", 2, "Presença", "Uso do Celular", 70.00),
    ("Maria José Junqueira da Silva", "19/09", "Sábado", "", 2, "Presença", "Uso do Celular", 70.00),
    ("Maria José Junqueira da Silva", "26/09", "Sábado", "", 2, "Presença", "Uso do Celular", 70.00),
    # Joelma
    ("Joelma Faria", "11/09", "Sexta", "", 1, "Falta", "Mkt Dig / AI / Ed Vídeos", 35.00),
    ("Joelma Faria", "17/09", "Quinta", "", 2, "Presença", "Mkt Dig / AI / Ed Vídeos", 70.00),
    # Maria Alice
    ("Maria Alice", "30/09", "Quarta", "", 2, "Presença", "Tráfego Pago", 70.00),
    # TURMA KIDS 1
    ("TURMA KIDS 1", "05/09", "Sábado", "15:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA KIDS 1", "12/09", "Sábado", "15:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA KIDS 1", "19/09", "Sábado", "15:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA KIDS 1", "26/09", "Sábado", "15:00", 2, "Presença", "Aula regular", 70.00),
    # TURMA ADOLESCENTES 1
    ("TURMA ADOLESCENTES 1", "05/09", "Sábado", "17:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA ADOLESCENTES 1", "12/09", "Sábado", "17:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA ADOLESCENTES 1", "19/09", "Sábado", "17:00", 2, "Presença", "Aula regular", 70.00),
    ("TURMA ADOLESCENTES 1", "26/09", "Sábado", "17:00", 2, "Presença", "Aula regular", 70.00),
    # TURMA KIDS 2
    ("TURMA KIDS 2", "29/09", "Terça", "17:00", 2, "Presença", "Aula regular", 70.00),
]

# Agrupa por aluno
alunos_ordem = []
alunos_dados = {}
for d in dados:
    if d[0] not in alunos_dados:
        alunos_ordem.append(d[0])
        alunos_dados[d[0]] = []
    alunos_dados[d[0]].append(d)

zebra = False
for aluno_nome in alunos_ordem:
    aulas_aluno = alunos_dados[aluno_nome]
    subtotal = sum(a[7] for a in aulas_aluno)

    # Header do aluno
    ws.merge_cells(f'A{linha}:H{linha}')
    ws[f'A{linha}'] = f"📋 {aluno_nome}"
    ws[f'A{linha}'].font = Font(name='Calibri', size=12, bold=True, color=VERDE_ESCURO)
    ws[f'A{linha}'].fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
    ws[f'A{linha}'].alignment = align_left
    ws.row_dimensions[linha].height = 24
    linha += 1

    # Cabeçalho das colunas
    headers = ["#", "Aluno", "Data", "Dia", "Horário", "Duração", "Conteúdo", "Valor"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=linha, column=col, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = borda
    ws.row_dimensions[linha].height = 22
    linha += 1

    # Aulas
    for idx, a in enumerate(aulas_aluno, 1):
        valores = [idx, a[0], a[1], a[2], a[4] if a[4] else "—", f"{a[4]}h" if a[4] else "—", a[6], a[7]]
        for col, v in enumerate(valores, 1):
            cell = ws.cell(row=linha, column=col, value=v)
            cell.font = font_normal
            cell.border = borda
            if col == 8:
                cell.number_format = 'R$ #,##0.00'
                cell.alignment = align_right
            elif col in (1, 3, 4, 5, 6):
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            # Zebra
            if zebra:
                cell.fill = fill_zebra
            # Falta em vermelho
            if a[5] == "Falta":
                cell.fill = fill_falta
                if col == 6:
                    cell.font = Font(name='Calibri', size=11, bold=True, color=VERMELHO)
        ws.row_dimensions[linha].height = 20
        linha += 1
        zebra = not zebra

    # Subtotal do aluno
    ws.merge_cells(f'A{linha}:G{linha}')
    ws[f'A{linha}'] = f"Subtotal {aluno_nome}"
    ws[f'A{linha}'].font = font_subtotal
    ws[f'A{linha}'].fill = fill_subtotal
    ws[f'A{linha}'].alignment = align_right
    ws[f'A{linha}'].border = borda
    cell_val = ws.cell(row=linha, column=8, value=subtotal)
    cell_val.font = font_subtotal
    cell_val.fill = fill_subtotal
    cell_val.number_format = 'R$ #,##0.00'
    cell_val.alignment = align_right
    cell_val.border = borda
    ws.row_dimensions[linha].height = 24
    linha += 1
    zebra = False
    linha += 1  # linha em branco entre alunos

# ===== TOTAL GERAL =====
linha += 1
ws.merge_cells(f'A{linha}:G{linha}')
ws[f'A{linha}'] = "TOTAL GERAL — SETEMBRO 2026"
ws[f'A{linha}'].font = Font(name='Calibri', size=14, bold=True, color=DOURADO)
ws[f'A{linha}'].fill = fill_total
ws[f'A{linha}'].alignment = align_center
ws[f'A{linha}'].border = Border(left=borda_titulo, right=borda_titulo, top=borda_titulo, bottom=borda_titulo)
total_geral = sum(d[7] for d in dados)
cell_total = ws.cell(row=linha, column=8, value=total_geral)
cell_total.font = Font(name='Calibri', size=14, bold=True, color=DOURADO)
cell_total.fill = fill_total
cell_total.number_format = 'R$ #,##0.00'
cell_total.alignment = align_center
cell_total.border = Border(left=borda_titulo, right=borda_titulo, top=borda_titulo, bottom=borda_titulo)
ws.row_dimensions[linha].height = 40
linha += 2

# ===== RESUMO =====
ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = "📊 RESUMO"
ws[f'A{linha}'].font = Font(name='Calibri', size=12, bold=True, color=VERDE_ESCURO)
ws[f'A{linha}'].alignment = align_left
linha += 1

resumo = [
    ("Total de aulas:", len(dados)),
    ("Presenças:", sum(1 for d in dados if d[5] == "Presença")),
    ("Faltas:", sum(1 for d in dados if d[5] == "Falta")),
    ("Horas trabalhadas:", f"{sum(d[4] for d in dados if d[5] == 'Presença') + sum(1 for d in dados if d[5] == 'Falta')}h"),
    ("Total a receber:", f"R$ {total_geral:.2f}".replace('.', ',')),
]
for label, valor in resumo:
    ws.cell(row=linha, column=2, value=label).font = Font(name='Calibri', size=11, bold=True, color='374151')
    ws.cell(row=linha, column=2).alignment = align_left
    ws.cell(row=linha, column=4, value=valor).font = Font(name='Calibri', size=11, color=VERDE)
    ws.cell(row=linha, column=4).alignment = align_left
    linha += 1

linha += 1
ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = "Relatório gerado pelo Controle de Aulas — Sistema de gestão para professores"
ws[f'A{linha}'].font = font_rodape
ws[f'A{linha}'].alignment = align_center
linha += 1
ws.merge_cells(f'A{linha}:H{linha}')
ws[f'A{linha}'] = f"Gerado em: {os.popen('date').read().strip()}"
ws[f'A{linha}'].font = font_rodape
ws[f'A{linha}'].alignment = align_center

# ===== METADADOS =====
wb.properties.creator = "Controle de Aulas"
wb.properties.title = "Fechamento Setembro 2026 — Guilherme Miranda"

# ===== SALVAR =====
output_path = "/home/z/my-project/download/fechamento-setembro-2026.xlsx"
wb.save(output_path)
print(f"✅ Excel salvo: {output_path}")
print(f"   {len(dados)} aulas | {len(alunos_ordem)} alunos | R$ {total_geral:.2f}")
